"""Customer-specific document requirements (the "Accounts / Customers Extra Requirements" workbook).

Two halves:
  1. parse_workbook(): every per-pod sheet with the standard header (Pod/Team, AM, Customer Name, BOL, POD, Notes)
     becomes one rule record per customer. The BOL/POD booleans are taken as-is; the free-text Notes are turned
     into typed flags by keyword rules (first pass, meant for AM review), and the full note is kept verbatim.
     The DetentionLayovers sheet sets deliver_out_with_detention; SOP sheets are kept as global text.
  2. check_document(): evaluates a filed document (extraction + type + load context) against the customer's
     rules and returns pass / fail / unknown per rule plus a deliver-out readiness verdict.

Nothing here changes TransportPro. The verdict is advisory: it goes into the File History comment, the dispatch
note, and the review queue, so a person sees "POD received but missing receiver signature (Win-Holt requires
driver + receiver)" instead of discovering it at billing.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

from .index import norm_name
from .schema import Extraction

STANDARD_HEADER = ("pod/team", "am", "customer name", "bol", "pod", "notes")
GLOBAL_SHEETS = ("detention requests- fischer pod", "matthwew gardner updated sop")
DETENTION_SHEET = "detentionlayovers"


# ---------------------------------------------------------------- parsing ----

def _flag(text: str, pattern: str) -> bool:
    return re.search(pattern, text, re.I) is not None


def rules_from_note(note: str, bol_required: bool, pod_required: bool) -> dict:
    """Keyword first pass over a free-text SOP note. Conservative: a flag is set only on clear wording."""
    t = " ".join(note.split())
    pages: int | str | None = None
    m = re.search(r"all\s+(\d+)\s+pages|\((\d+)\s+pages\)|(\d+)\s+pages\s+of\s+(?:the\s+)?bol", t, re.I)
    if m:
        pages = int(next(g for g in m.groups() if g))
    elif _flag(t, r"\bboth pages\b|\btwo total\b"):
        pages = 2
    elif _flag(t, r"all (of the )?(stated |necessary )?pages|all pages of the bol|how many pages"):
        pages = "per_load_note" if _flag(t, r"load notes|load info|stated pages") else "all"

    pod_sigs: list[str] = []
    if pod_required or _flag(t, r"\bpod\b"):
        if _flag(t, r"receiver'?s? signature|signed pod|pod .{0,40}signed|receiver.{0,30}signature|both signatures|driver & receiver|driver and receiver|shipper, receiver, and driver"):
            pod_sigs.append("receiver")
        if _flag(t, r"driver & receiver|driver and receiver|both signatures|shipper, receiver, and driver"):
            pod_sigs.append("driver")
    bol_sigs: list[str] = []
    if _flag(t, r"bol needs signed|signed copy of bol|signed bol|bol must have shipper and driver signatures"):
        bol_sigs = ["shipper", "driver"] if _flag(t, r"shipper and driver") else ["shipper"]

    deliver_out_with_detention: bool | None = None
    if _flag(t, r"can be delivered out pending detention|good to deliver out pending detention|regardless of detention|dont need to wait to deliver out|do not have to wait for paperwork"):
        deliver_out_with_detention = True
    elif _flag(t, r"leave open pending detention|do not deliver out the load|not deliver out"):
        deliver_out_with_detention = False

    doc_type_overrides: dict[str, str] = {}
    if _flag(t, r"mark the pod from the receiver as the bill of lading"):
        doc_type_overrides["receiver_signed"] = "Bill Of Lading"
    if _flag(t, r"driver supplied bol"):
        doc_type_overrides["shipper_copy"] = "Driver Supplied BOL"

    return {
        "bol_before_leaving_shipper": bol_required and _flag(t, r"before (leav|depart)|prior to leav|before setting dispatch status to .?loaded|once loaded|immediately needed upon being loaded|before they leave"),
        "pod_before_deliver_out": pod_required and _flag(t, r"deliver(ing|ed)? (this |the load |it )?out|set (as|to) delivered|close out the load|to deliver out"),
        "upload_required_before_deliver_out": _flag(t, r"upload"),
        "pod_within_minutes": int(m2.group(1)) if (m2 := re.search(r"pod (?:with)?in (\d+) ?min", t, re.I)) else (24 * 60 if _flag(t, r"pod required within 24 ?hrs") else None),
        "pages_required": pages,
        "bol_signatures": bol_sigs,
        "pod_signatures": pod_sigs,
        "stamp_acceptable": _flag(t, r"\bstamp\b"),
        "seal_required_on_bol": bol_required and _flag(t, r"\bseal\b") and not _flag(t, r"seal pic is not required"),
        "freight_photos_required": _flag(t, r"picture[s]? of (the )?(loaded )?(freight|trailer|load)|pics of freight|photos of the load"),
        "in_out_times_required": _flag(t, r"in/out|in and out times|in & out|with date and time|times must be"),
        "address_match_required": _flag(t, r"matches the load|right address|correct freight/ppw|check the crate count"),
        "escalate_if_missing": _flag(t, r"escalat") and not _flag(t, r"do not need to escalate|dont need to escalate"),
        "deliver_out_with_detention": deliver_out_with_detention,
        "doc_type_overrides": doc_type_overrides,
        "macropoint_required": _flag(t, r"macropoint|\bMP\b.{0,20}(non-?negotiable|required|not negotiable)"),
        "customer_email_updates": _flag(t, r"email chain|custy email|customer notification|portal update"),
    }


def load_pod_map(path: str | Path | None) -> dict:
    """pod_terminals.json: TransportPro terminal IDs per pod and the sheet -> terminal mapping."""
    if not path or not Path(path).exists():
        return {"terminals": [], "sheet_map": [], "pods_without_sheet": []}
    return json.loads(Path(path).read_text(encoding="utf-8"))


def parse_workbook(path: str | Path, pod_map_path: str | Path | None = None) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    customers: list[dict] = []
    globals_: dict[str, str] = {}
    detention: dict[str, dict] = {}
    pod_map = load_pod_map(pod_map_path)
    sheet_terminals = {m["sheet"].strip().lower(): m for m in pod_map.get("sheet_map", [])}

    for ws in wb.worksheets:
        title = ws.title.strip().lower()
        if title in GLOBAL_SHEETS:
            globals_[ws.title.strip()] = "\n".join(" | ".join(str(v).strip() for v in r if v is not None and str(v).strip()) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r))
            continue
        if title == DETENTION_SHEET:
            for r in ws.iter_rows(values_only=True):
                vals = [("" if v is None else str(v).strip()) for v in r]
                if len(vals) >= 3 and vals[0] and vals[0] != "Customer" and vals[2] in ("True", "False"):
                    detention[norm_name(vals[0])] = {"customer": vals[0], "pod": vals[1], "deliver_out": vals[2] == "True", "note": vals[3] if len(vals) > 3 else ""}
            continue
        header_row = None
        for i, r in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
            cells = tuple((str(v).strip().lower() if v is not None else "") for v in r[:6])
            if cells == STANDARD_HEADER:
                header_row = i
                break
        if header_row is None:
            continue
        for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
            vals = [("" if v is None else str(v).strip()) for v in r]
            if len(vals) < 6 or not vals[2] or vals[3] not in ("True", "False"):
                continue
            note = vals[5]
            extra = " || ".join(v for v in vals[6:] if v)
            bol_req, pod_req = vals[3] == "True", vals[4] == "True"
            sm = sheet_terminals.get(ws.title.strip().lower(), {})
            rec = {
                "customer": vals[2], "customer_key": norm_name(vals[2]), "customer_base": norm_name(re.split(r"\bc/o\b|\(", vals[2], flags=re.I)[0]),
                "sheet": ws.title.strip(), "pod_team": vals[0], "am": vals[1],
                "terminal_ids": sm.get("terminal_ids", []), "terminal_map_confidence": sm.get("confidence", "unknown"),
                "bol_required": bol_req, "pod_required": pod_req,
                **rules_from_note(note + " " + extra, bol_req, pod_req),
                "source_note": note, "source_extra": extra,
            }
            customers.append(rec)

    for rec in customers:                         # merge the detention sheet by name
        for key, d in detention.items():
            if key and (key in rec["customer_key"] or rec["customer_key"] in key or key in rec["customer_base"]):
                if rec["deliver_out_with_detention"] is None:
                    rec["deliver_out_with_detention"] = d["deliver_out"]
    return {"_meta": {"source": Path(path).name, "customers": len(customers), "pod_map_source": str(pod_map_path) if pod_map_path else None},
            "customers": customers, "detention_layovers": list(detention.values()), "global_sops": globals_,
            "pods": pod_map}


# ---------------------------------------------------------------- lookup -----

class Requirements:
    def __init__(self, data: dict):
        self.data = data
        self.customers = data["customers"]
        pods = data.get("pods") or {}
        self.terminals = {t["id"]: t for t in pods.get("terminals", [])}
        self.non_pod_terminals = {t["id"]: t for t in pods.get("non_pod_terminals", [])}
        self.sheet_map = pods.get("sheet_map", [])
        self.pods_without_sheet = {p["id"]: p for p in pods.get("pods_without_sheet", [])}

    @classmethod
    def from_file(cls, path: str | Path) -> "Requirements":
        return cls(json.loads(Path(path).read_text(encoding="utf-8")))

    # ---- pods / terminals ----
    def pod_name(self, terminal_id: int | None) -> str | None:
        t = self.terminals.get(terminal_id) if terminal_id else None
        if t:
            return t["name"]
        t = self.non_pod_terminals.get(terminal_id) if terminal_id else None
        return f"{t['name']} (not a pod)" if t else None

    def sheet_for_terminal(self, terminal_id: int | None) -> dict | None:
        """The requirements sheet that covers a TransportPro terminal, or None."""
        if not terminal_id:
            return None
        return next((m for m in self.sheet_map if terminal_id in m.get("terminal_ids", [])), None)

    def pod_context(self, terminal_id: int | None) -> dict:
        """What the pipeline should say about the pod: name, sheet, and whether rules exist for it."""
        sheet = self.sheet_for_terminal(terminal_id)
        return {
            "terminal_id": terminal_id,
            "pod": self.pod_name(terminal_id),
            "sheet": sheet["sheet"] if sheet else None,
            "map_confidence": sheet["confidence"] if sheet else None,
            "status": ("rules on file" if sheet else ("pod known, no requirements sheet" if terminal_id in self.pods_without_sheet or terminal_id in self.terminals else "terminal not a Track & Trace pod")) if terminal_id else "no terminal on load",
        }

    # ---- customers ----
    def _score(self, rec: dict, n: str, base: str) -> float:
        score = 0.0
        for a in (rec["customer_key"], rec["customer_base"]):
            for b in (n, base):
                if not a or not b:
                    continue
                s = 1.0 if a == b else (0.9 if (a in b or b in a) and min(len(a), len(b)) >= 5 else SequenceMatcher(None, a, b).ratio())
                score = max(score, s)
        return score

    def for_customer(self, name: str | None, terminal_id: int | None = None) -> dict | None:
        """Fuzzy lookup by TransportPro customer name; tolerates 'c/o' suffixes and Inc/LLC noise.

        When the load's terminal is known, rows from that pod's own sheet win; a row from another pod's sheet is
        returned only if no own-pod row matches, and it is tagged cross_pod=True so the verdict can say so.
        """
        if not name:
            return None
        n = norm_name(name)
        base = norm_name(re.split(r"\bc/o\b|\(", name, flags=re.I)[0])
        own_sheet = self.sheet_for_terminal(terminal_id)
        own, other = (None, 0.0), (None, 0.0)
        for rec in self.customers:
            s = self._score(rec, n, base)
            if s < 0.8:
                continue
            if own_sheet and rec["sheet"].strip().lower() == own_sheet["sheet"].strip().lower():
                if s > own[1]:
                    own = (rec, s)
            elif s > other[1]:
                other = (rec, s)
        if own[0]:
            return {**own[0], "cross_pod": False}
        if other[0]:
            return {**other[0], "cross_pod": bool(terminal_id)}
        return None


# ---------------------------------------------------------------- checking ---

@dataclass
class RuleResult:
    rule: str
    status: str            # pass | fail | unknown | n/a
    detail: str


@dataclass
class Verdict:
    customer: str | None
    results: list[RuleResult] = field(default_factory=list)

    @property
    def failed(self) -> list[RuleResult]:
        return [r for r in self.results if r.status == "fail"]

    @property
    def unknown(self) -> list[RuleResult]:
        return [r for r in self.results if r.status == "unknown"]

    @property
    def summary(self) -> str:
        if self.customer is None:
            return "no customer-specific requirements on file; default rules apply"
        if self.failed:
            return "does NOT meet " + self.customer + " requirements: " + "; ".join(r.detail for r in self.failed)
        if self.unknown:
            return "meets checkable " + self.customer + " requirements; unverified: " + "; ".join(r.detail for r in self.unknown)
        return "meets " + self.customer + " requirements"


def check_document(ex: Extraction, filed_type: str, rules: dict | None, load: dict | None, name_signals: set[str]) -> Verdict:
    """filed_type is the TransportPro type the pipeline chose ('Bill Of Lading', 'Proof of Delivery', ...).
    name_signals: matcher signal names that fired (e.g. 'shipper_name', 'consignee_name') for the address check."""
    v = Verdict(customer=rules["customer"] if rules else None)
    if not rules:
        return v
    is_pod = filed_type == "Proof of Delivery" or (filed_type == "Bill Of Lading" and rules.get("doc_type_overrides", {}).get("receiver_signed") and ex.signatures.receiver_signed)
    is_bol = filed_type in ("Bill Of Lading", "Driver Supplied BOL") and not is_pod
    add = lambda rule, status, detail: v.results.append(RuleResult(rule, status, detail))

    # pages
    pr = rules.get("pages_required")
    if pr:
        n = len(ex.pages)
        if isinstance(pr, int):
            add("pages", "pass" if n >= pr else "fail", f"{n} of {pr} required pages present")
        elif pr == "per_load_note":
            add("pages", "unknown", f"{n} page(s) received; required count is in the load notes, confirm")
        else:
            add("pages", "unknown", f"{n} page(s) received; customer requires all pages")

    # signatures
    if is_pod:
        req = rules.get("pod_signatures") or []
        if "receiver" in req or rules.get("pod_required"):
            ok = ex.signatures.receiver_signed or (rules.get("stamp_acceptable") and ex.signatures.stamp_present)
            add("pod_receiver_signature", "pass" if ok else "fail", "receiver signature" + (" or stamp" if rules.get("stamp_acceptable") else "") + (" present" if ok else " missing"))
        if "driver" in req:
            add("pod_driver_signature", "pass" if ex.signatures.driver_signed else "fail", "driver signature " + ("present" if ex.signatures.driver_signed else "missing"))
    if is_bol:
        for who in rules.get("bol_signatures") or []:
            signed = getattr(ex.signatures, f"{who}_signed", False)
            add(f"bol_{who}_signature", "pass" if signed else "fail", f"{who} signature " + ("present" if signed else "missing"))
        if rules.get("seal_required_on_bol"):
            has_seal = any(n.kind == "seal" for n in ex.numbers)
            add("seal_on_bol", "pass" if has_seal else "fail", "seal number " + ("present" if has_seal else "not found on BOL"))

    # in/out times
    if rules.get("in_out_times_required") and is_pod:
        # A customer that requires in/out times wants the consignee's. Pickup times satisfy nothing,
        # and before at_stop existed they were indistinguishable. As in classify_type, only an
        # explicit "shipper" fails: "unknown" keeps the behaviour every earlier extraction was judged by.
        at_shipper = ex.times.at_stop == "shipper"
        both = bool(ex.times.check_in and ex.times.check_out)
        ok = both and not at_shipper
        where = f", at the {ex.times.at_stop}" if ex.times.at_stop != "unknown" else ""
        shown = f"{ex.times.check_in} / {ex.times.check_out} ({ex.times.source}{where})"
        why = shown if ok else (f"{shown} - recorded at the shipper, not the consignee" if both else "missing on POD")
        add("in_out_times", "pass" if ok else "fail", "in/out times " + why)

    # address / customer on paper matches the load
    if rules.get("address_match_required"):
        ok = bool({"shipper_name", "consignee_name", "customer_name"} & name_signals)
        add("address_match", "pass" if ok else "unknown", "paper parties match the load" if ok else "could not confirm paper parties against the load")

    # freight photos: only checkable if a photo page arrived in the same submission
    if rules.get("freight_photos_required") and is_bol:
        has_photo = any(p.role == "photo" for p in ex.pages)
        add("freight_photos", "pass" if has_photo else "unknown", "freight photo included" if has_photo else "freight photos required before leaving shipper; none in this submission")

    # timing relative to dispatch stage (informational)
    stage = (load or {}).get("dispatch_status", "")
    if is_bol and rules.get("bol_before_leaving_shipper") and stage in ("Delivered", "At Consignee"):
        add("bol_timing", "fail", f"BOL required before leaving the shipper; arrived with dispatch {stage}")

    # deliver-out readiness
    if rules.get("pod_required"):
        ready = is_pod and not [r for r in v.results if r.status == "fail"]
        add("deliver_out_ready", "pass" if ready else "fail", "POD on file and passes checks; load may be delivered out" if ready else "POD required before delivering out; not yet satisfied")
    if rules.get("deliver_out_with_detention") is False:
        add("detention_hold", "unknown", "customer requires the load to stay open while a detention/layover request is pending")
    return v
